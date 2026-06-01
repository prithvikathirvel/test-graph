import io
import json
import base64
import asyncio
import logging
from typing import Optional
from pydantic import BaseModel, Field

# --- All necessary imports, including the new tool converter ---
from app.core.config import settings
from app.engine.registry import NodeRegistry
from app.utils.templating import resolve_placeholders
from app.core.state import FlowState
from app.core.token_tracker import token_callback_handler

from langchain_core.language_models import BaseChatModel
from langchain_core.prompts import ChatPromptTemplate, MessagesPlaceholder
from langchain_core.messages import HumanMessage, AIMessage
from langchain_core.utils.function_calling import convert_to_openai_tool
from langchain_google_vertexai import ChatVertexAI
from langchain_openai import ChatOpenAI
from google.oauth2 import service_account

logger = logging.getLogger(__name__)


# ── 1. Strictly Typed Chatty Output (Unchanged) ────────────────────────────────
class ChattyResponse(BaseModel):
    message: str = Field(description="The actual text response to show to the user. Use HTML tags for formatting (<b>, <i>, <ul>, <li>) for better UI rendering. Use /n for formatting. Render with white-space: pre-line /n")
    intent: str = Field(description="MUST be one of: 'continue', 'close', 'escalate', 'tool_call'.")
    tool_name: Optional[str] = Field(None, description="If intent is 'tool_call', the exact name of the tool to execute.")


# ── 2. Model Name Normalizer for Gemini (Unchanged) ───────────────────────────
_MODEL_NAME_MAP = {
    # ... (your gemini models)
    "gemini 2.5 pro": "gemini-2.5-pro",
    "gemini 1.5 pro": "gemini-1.5-pro-002",
}

def _normalize_model_name(raw: str) -> str:
    key = raw.strip().lower()
    if key in _MODEL_NAME_MAP: return _MODEL_NAME_MAP[key]
    if "-" in key and " " not in key: return raw.strip()
    fallback = key.replace(" ", "-")
    logger.warning(f"Unknown Gemini model display name '{raw}'. Falling back to '{fallback}'.")
    return fallback


# ── 3. CPU-Bound Document Converters (Unchanged) ───────────────────────────────────
def _convert_to_pdf_b64(text: str) -> str:
    try:
        from fpdf import FPDF
    except ImportError:
        raise ImportError("Please install fpdf2: pip install fpdf2")
    pdf = FPDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.set_font("Helvetica", size=12)
    clean_text = text.encode('latin-1', 'replace').decode('latin-1')
    pdf.multi_cell(0, 8, clean_text)
    return base64.b64encode(pdf.output()).decode('utf-8')

def _convert_to_docx_b64(text: str) -> str:
    try:
        from docx import Document
    except ImportError:
        raise ImportError("Please install python-docx: pip install python-docx")
    doc = Document()
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return base64.b64encode(buf.getvalue()).decode('utf-8')


# ── 4. LLM Model Factory (Unchanged from previous refactor) ──────────────────
try:
    credentials = service_account.Credentials.from_service_account_file(
        "service-account.json",
        scopes=["https://www.googleapis.com/auth/cloud-platform"],
    )
except Exception as e:
    logger.warning(f"Could not load GCP credentials: {e}")
    credentials = None

def _get_llm(model_name: str, temperature: float = 0.0) -> BaseChatModel:
    normalized_name = model_name.strip().lower()
    
    # ── 1. Llama / InfinitAI Maas Path ───────────────────────────────────────
    if "llama" in normalized_name:
        # If the user provided a full path (e.g. "meta-llama/Llama-3-70b-instruct"), use it.
        # Otherwise, map generic names to the best available version on the provider.
        model_id = model_name.strip()
        if normalized_name in ["llama 3", "llama3", "llama 3.1", "llama3.1", "llama 3.3", "llama3.3"]:
            # model_id = "meta-llama/Llama-3.1-8B-Instruct"
            model_id = "meta/llama-3.3-70b-instruct"
        
        logger.info(f"Initializing ChatOpenAI for Llama model: '{model_id}' via {settings.OPENAI_BASE_URL}")
        return ChatOpenAI(
            model=model_id,
            api_key=settings.OPENAI_API_KEY,
            base_url=settings.OPENAI_BASE_URL,
            temperature=temperature,
            callbacks=[token_callback_handler],
        )
    
    # ── 2. Gemini / VertexAI Path ────────────────────────────────────────────
    else:
        slug = _normalize_model_name(model_name)
        logger.info(f"Initializing ChatVertexAI with Gemini model='{slug}'")
        return ChatVertexAI(
            model=slug,
            credentials=credentials,
            temperature=temperature,
            callbacks=[token_callback_handler],
        )


# ── 5. The Node Executor (Final Production Version) ────────────────────────
@NodeRegistry.register("LLM invoker")
async def llm_invoker_node(state: FlowState, node_config: dict) -> dict:
    logger.info("🟢 Entering LLM Invoker Node...")

    # 1. Parse configuration
    inputs = {p["key"]: p["value"] for p in node_config.get("inputParameters", [])}
    model_choice = inputs.get("Model", "gemini-2.5-flash")
    chatty_mode = str(inputs.get("chatty_mode", "false")).lower() == "true"
    user_msg_key = inputs.get("user_message_key", "CHAT_QUERY")
    memory_window = int(inputs.get("memory_window", 20))
    use_memory = bool(inputs.get("use_memory", True))
    response_format = str(inputs.get("Response Format", "text")).lower()
    out_params = node_config.get("outputParameters", [])
    output_key = out_params[0]["value"] if out_params else "response"

    # 2. Resolve dynamic variables
    raw_prompt = inputs.get("Prompt", "You are a helpful assistant.")
    system_prompt = resolve_placeholders(raw_prompt, state["variables"])
    user_query = str(state["variables"].get(user_msg_key, "")).strip()

    # CRITICAL: Gemini (VertexAI) fails with 400 if user_query is empty ("at least one parts field")
    if not user_query:
        logger.warning(f"User query from key '{user_msg_key}' was empty. Using fallback.")
        user_query = "..."  # Minimal non-empty string to avoid API failure

    if not system_prompt:
        system_prompt = "You are a helpful assistant."

    logger.info(f"Resolved System Prompt (len={len(system_prompt)})")
    logger.info(f"User Query: '{user_query[:200]}'")

    # 3. Trim history
    all_messages = state.get("messages", [])
    if use_memory and memory_window > 0:
        recent_chat_history = all_messages[-(memory_window * 2):]
    else:
        recent_chat_history = []
    
    logger.debug(f"Context: {len(recent_chat_history)} messages in history window (Memory: {'ON' if use_memory else 'OFF'}).")

    # 4. Initialize LLM
    try:
        llm = _get_llm(model_choice)
    except Exception as e:
        logger.error(f"❌ Failed to initialize LLM: {e}", exc_info=True)
        return {"variables": {output_key: {"message": f"Initialization Error: {e}", "intent": "continue"}}}

    # 4b. Prepare instructions based on format
    if response_format == "json" and not chatty_mode:
        system_prompt += "\n\nCRITICAL INSTRUCTION: You must respond with raw, valid JSON only. Do not wrap it in markdown blocks."
    elif response_format in ["text", "html"]:
        system_prompt += "\n\nUI RENDERING INSTRUCTION: Use HTML tags (<b>, <i>, <ul>, <li>, <br>) for formatting."

    # 5. Execute
    try:
        if chatty_mode:
            # ── CHATTY MODE (Final Multi-Provider Logic) ───
            tool_context = [f"[{k}]: {v}" for k, v in state["variables"].items() if k.endswith(("_output", "Status")) or k in ("context", "kb_context")]
            if tool_context:
                system_prompt += "\n\n### TOOL RESULTS (DO NOT CALL THESE TOOLS AGAIN) ###\n" + "\n".join(tool_context)

            prompt_template = ChatPromptTemplate.from_messages([
                ("system", "{system_prompt_var}"),
                MessagesPlaceholder(variable_name="chat_history"),
                ("human", "{user_query}"),
            ])
            
            response_obj: ChattyResponse

            # ✅ --- RECOMMENDED: Llama path using Manual Tool-Calling ---
            # This is the most robust method when .with_structured_output() is not supported by the endpoint.
            # It avoids the 400 Bad Request error by using the API's native function/tool calling mechanism.
            if isinstance(llm, ChatOpenAI):
                logger.debug("Using Llama (Manual Tool-Calling) path for structured output.")
                
                # 1. Convert the Pydantic model into the format the API expects for a tool.
                chatty_response_tool = convert_to_openai_tool(ChattyResponse)
                
                # 2. Bind the tool to the LLM and *force* the LLM to use it.
                llm_with_tool = llm.bind(
                    tools=[chatty_response_tool],
                    tool_choice={"type": "function", "function": {"name": "ChattyResponse"}}
                )
                
                chain = prompt_template | llm_with_tool
                
                # 3. Invoke the chain. The response will now contain a `tool_calls` attribute.
                response_message = await chain.ainvoke({
                    "system_prompt_var": system_prompt,
                    "chat_history": recent_chat_history,
                    "user_query": user_query,
                })
                
                # 4. Reliably parse the arguments from the tool call.
                if response_message.tool_calls:
                    tool_call = response_message.tool_calls[0]
                    response_obj = ChattyResponse(**tool_call['args'])
                else:
                    # Fallback if the model fails to use the tool
                    logger.warning("Llama model did not use the required tool. Falling back to text response.")
                    response_obj = ChattyResponse(message=response_message.content, intent="continue")

            # --- Gemini Path (Backward Compatible) ---
            elif isinstance(llm, ChatVertexAI):
                logger.debug("Using Gemini (.with_structured_output) path for structured output.")
                chain = prompt_template | llm.with_structured_output(ChattyResponse)
                response_obj = await chain.ainvoke({
                    "system_prompt_var": system_prompt,
                    "chat_history": recent_chat_history,
                    "user_query": user_query,
                })
            
            else:
                raise TypeError(f"Unsupported LLM type for chatty mode: {type(llm).__name__}")
            
            # This part is the same as it works with the parsed response_obj
            logger.info(f"LLM Response Intention: {response_obj.intent}")
            result_dict = {
                "message": response_obj.message,
                "intent": response_obj.intent if response_obj.intent in ["continue", "close", "escalate", "tool_call"] else "continue",
                "is_chatty": True,
                "user_text": user_query,
            }
            if response_obj.tool_name:
                result_dict["tool_name"] = response_obj.tool_name

            return {
                "variables": {output_key: result_dict},
                "messages": [HumanMessage(content=user_query), AIMessage(content=response_obj.message)],
            }

        else:
            # ── STANDARD MODE (Works for both models) ───────────
            prompt = ChatPromptTemplate.from_messages([
                ("system", "{system_prompt_var}"),
                MessagesPlaceholder(variable_name="chat_history"),
                ("human", "{user_query}"),
            ])
            chain = prompt | llm
            logger.debug(f"Executing standard LLM call (Format: {response_format})...")
            response = await chain.ainvoke({
                "system_prompt_var": system_prompt,
                "chat_history": recent_chat_history,
                "user_query": user_query,
            })

            final_text = response.content
            logger.debug(f"Received LLM response (len={len(final_text)})")
            new_vars: dict = {}

            try:
                if response_format == "pdf":
                    encoded = await asyncio.to_thread(_convert_to_pdf_b64, final_text)
                    new_vars[output_key] = encoded
                    new_vars["fileType"] = "pdf"
                    new_vars["fileName"] = "output.pdf"

                elif response_format in ["doc", "docx", "document", "docs"]:
                    encoded = await asyncio.to_thread(_convert_to_docx_b64, final_text)
                    new_vars[output_key] = encoded
                    new_vars["fileType"] = "doc"
                    new_vars["fileName"] = "output.docx"

                elif response_format == "json":
                    clean_text = final_text.strip()
                    if clean_text.startswith("```json"):
                        clean_text = clean_text[7:]
                    elif clean_text.startswith("```"):
                        clean_text = clean_text[3:]
                    if clean_text.endswith("```"):
                        clean_text = clean_text[:-3]
                    clean_text = clean_text.strip()

                    try:
                        new_vars[output_key] = json.loads(clean_text)
                    except json.JSONDecodeError as je:
                        logger.error(f"LLM failed to produce valid JSON: {je}. Raw: {clean_text}")
                        new_vars[output_key] = {"error": "LLM did not return valid JSON", "raw_response": clean_text}

                else:
                    new_vars[output_key] = final_text

            except Exception as format_err:
                logger.error(f"Failed to convert format to {response_format}: {format_err}", exc_info=True)
                new_vars[output_key] = final_text  # safe fallback

            return {
                "variables": new_vars,
                "messages":  [HumanMessage(content=user_query), AIMessage(content=final_text)],
            }

    except Exception as e:
        logger.exception("❌ LLM Invoker Error")
        error_msg = f"I'm sorry, an error occurred: {str(e)}"
        if chatty_mode:
            return {"variables": {output_key: {"message": error_msg, "intent": "continue", "error": str(e)}}}
        return {"variables": {output_key: error_msg}}